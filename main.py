from fastapi import FastAPI, UploadFile, File
from fastapi.responses import JSONResponse
import tempfile
import os
import subprocess
import json

import fitz  
import docx  
import openpyxl 


import tiktoken

app = FastAPI()



def extract_text_from_pdf(path: str) -> str:
    text = []
    doc = fitz.open(path)
    for page in doc:
        text.append(page.get_text())
    return "\n".join(text)


def extract_text_from_docx(path: str) -> str:
    doc = docx.Document(path)
    return "\n".join([para.text for para in doc.paragraphs])


def extract_text_from_xlsx(path: str) -> str:
    wb = openpyxl.load_workbook(path, data_only=True)
    text = []
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=True):
            for cell in row:
                if cell:
                    text.append(str(cell))
    return "\n".join(text)


def get_text_from_file(file_path: str, filename: str) -> str:
    ext = filename.lower().split(".")[-1]
    if ext == "pdf":
        return extract_text_from_pdf(file_path)
    if ext == "docx":
        return extract_text_from_docx(file_path)
    if ext in ("xlsx", "xls"):
        return extract_text_from_xlsx(file_path)
    # fallback to plain text
    with open(file_path, "r", errors="ignore") as f:
        return f.read()



def count_tokens_with_node(text: str) -> int:
    with tempfile.NamedTemporaryFile(delete=False, mode="w", encoding="utf-8") as f:
        f.write(text)
        tmp_path = f.name

    try:
        result = subprocess.run(
            ["node", "token_counter.js", tmp_path],
            capture_output=True,
            text=True
        )
        try:
            data = json.loads(result.stdout)
            return data.get("token_count", 0)
        except json.JSONDecodeError:
            print("Node output not JSON:", result.stdout)
            return 0
    finally:
        os.remove(tmp_path)



def count_tokens_with_tiktoken(text: str, model_name: str = "gpt-4") -> int:

    enc = tiktoken.encoding_for_model(model_name)
    return len(enc.encode(text))


@app.post("/count-tokens")
async def count_tokens(data: UploadFile = File(...)):

    try:
        file = data
        with tempfile.NamedTemporaryFile(delete=False) as tmp:
            content = await file.read()
            tmp.write(content)
            tmp_path = tmp.name

        text = get_text_from_file(tmp_path, file.filename)
        token_count = count_tokens_with_node(text)

        return JSONResponse({
            "filename": file.filename,
            "token_count": token_count
        })
    finally:
        os.remove(tmp_path)


@app.post("/count-tokens/tiktoken")
async def count_tokens_tiktoken(data: UploadFile = File(...), model: str = "gpt-4"):

    try:
        file = data
        with tempfile.NamedTemporaryFile(delete=False) as tmp:
            content = await file.read()
            tmp.write(content)
            tmp_path = tmp.name

        text = get_text_from_file(tmp_path, file.filename)
        token_count = count_tokens_with_tiktoken(text, model_name=model)

        return JSONResponse({
            "filename": file.filename,
            "model": model,
            "token_count": token_count
        })
    finally:
        os.remove(tmp_path)

@app.post("/count-tokens/tiktokenv2")
async def count_tokens_tiktoken(data: UploadFile = File(...), model: str = "gpt-4"):

    try:
        file = data
        with tempfile.NamedTemporaryFile(delete=False) as tmp:
            content = await file.read()
            tmp.write(content)
            tmp_path = tmp.name

        text = get_text_from_file(tmp_path, file.filename)
        token_count = count_tokens_with_tiktoken(text, model_name=model)

        return JSONResponse({
            "filename": file.filename,
            "model": "Proprietary Models",
            "token_count": token_count
        })
    finally:
        os.remove(tmp_path)

